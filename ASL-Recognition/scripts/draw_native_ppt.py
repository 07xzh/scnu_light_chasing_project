import pptx
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

prs = pptx.Presentation(r'C:\Users\20229\Desktop\ASL_Pro_Studio_Presentation_V2.pptx')

slide_arch = prs.slides[2]
# Clear all shapes on slide 2 except the title
for shape in list(slide_arch.shapes):
    if shape.has_text_frame and '系统架构' not in shape.text:
        sp = shape.element
        sp.getparent().remove(sp)
    elif shape.shape_type == 13: # Picture
        sp = shape.element
        sp.getparent().remove(sp)

def add_box(slide, x, y, w, h, text, fill_rgb=(255,255,255), line_rgb=(0,0,0), font_size=12, font_color=(0,0,0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_rgb)
    shape.line.color.rgb = RGBColor(*line_rgb)
    tf = shape.text_frame
    tf.text = text
    for paragraph in tf.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(*font_color)
            run.font.name = 'Microsoft YaHei'
    return shape

# Draw Frontend Container
add_box(slide_arch, 0.5, 1.5, 9, 1.5, "Frontend UI (index.html)", (240, 248, 255), (70, 130, 180), 14, (25, 25, 112))
c1 = add_box(slide_arch, 1, 2.1, 1.5, 0.6, "Canvas 渲染\n(60FPS)", (224, 255, 255), (0, 139, 139), 11)
c2 = add_box(slide_arch, 3, 2.1, 1.5, 0.6, "手势捕获面板\n(UI Layer)", (224, 255, 255), (0, 139, 139), 11)
c3 = add_box(slide_arch, 5, 2.1, 1.5, 0.6, "翻译显示区\n(Result)", (224, 255, 255), (0, 139, 139), 11)
c4 = add_box(slide_arch, 7, 2.1, 1.5, 0.6, "沉浸式侧边栏\n(Controls)", (224, 255, 255), (0, 139, 139), 11)

# Draw Backend Container
add_box(slide_arch, 0.5, 3.5, 9, 3.8, "Backend (FastAPI & Models)", (245, 255, 250), (46, 139, 87), 14, (0, 100, 0))
ws = add_box(slide_arch, 4.25, 4.1, 1.5, 0.6, "ws_endpoint\n(async WebSocket)", (255, 250, 205), (184, 134, 11), 11)
cam = add_box(slide_arch, 1, 4.1, 1.5, 0.6, "camera_loop\n(独立多线程 66fps)", (255, 250, 205), (184, 134, 11), 11)
snd = add_box(slide_arch, 7.5, 4.1, 1.5, 0.6, "send_loop\n(40Hz 高频推送)", (255, 250, 205), (184, 134, 11), 11)

pipe = add_box(slide_arch, 1, 5.3, 1.5, 0.6, "HandASLPipeline\n(MediaPipe+Keras)", (255, 228, 225), (205, 92, 92), 11)
sub = add_box(slide_arch, 4.25, 5.3, 1.5, 0.6, "submit_word()\n(核心业务逻辑)", (255, 228, 225), (205, 92, 92), 11)
state = add_box(slide_arch, 7.5, 5.3, 1.5, 0.6, "Global State\n(SSOT 单一数据源)", (230, 230, 250), (147, 112, 219), 11)

spel = add_box(slide_arch, 1.5, 6.5, 1.5, 0.6, "SpellChecker\n(字典纠错)", (255, 240, 245), (219, 112, 147), 11)
shor = add_box(slide_arch, 3.5, 6.5, 1.5, 0.6, "ShortcutManager\n(快捷词展开)", (255, 240, 245), (219, 112, 147), 11)
llm = add_box(slide_arch, 5.5, 6.5, 1.5, 0.6, "DeepSeek API\n(LLM 成句)", (255, 240, 245), (219, 112, 147), 11)
auth = add_box(slide_arch, 7.5, 6.5, 1.5, 0.6, "Face Auth\n(活体检测+认证)", (255, 240, 245), (219, 112, 147), 11)

from pptx.enum.shapes import MSO_CONNECTOR
def add_connector(slide, begin_x, begin_y, end_x, end_y):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(begin_x), Inches(begin_y), Inches(end_x), Inches(end_y))
    connector.line.color.rgb = RGBColor(100, 100, 100)
    connector.line.width = Pt(1.5)

add_connector(slide_arch, 1.75, 2.7, 5.0, 4.1)
add_connector(slide_arch, 3.75, 2.7, 5.0, 4.1)
add_connector(slide_arch, 5.75, 2.7, 5.0, 4.1)
add_connector(slide_arch, 7.75, 2.7, 5.0, 4.1)
add_connector(slide_arch, 1.75, 4.7, 1.75, 5.3)
add_connector(slide_arch, 5.0, 4.7, 5.0, 5.3)
add_connector(slide_arch, 8.25, 4.7, 8.25, 5.3)
add_connector(slide_arch, 2.5, 5.6, 4.25, 5.6)
add_connector(slide_arch, 5.75, 5.6, 7.5, 5.6)
add_connector(slide_arch, 4.75, 5.9, 2.25, 6.5)
add_connector(slide_arch, 5.0, 5.9, 4.25, 6.5)
add_connector(slide_arch, 5.25, 5.9, 6.25, 6.5)
add_connector(slide_arch, 8.25, 6.5, 8.25, 5.9)

# Slide 3 (Pipeline)
slide_pipe = prs.slides[3]
for shape in list(slide_pipe.shapes):
    if shape.has_text_frame and '核心识别流水线' not in shape.text:
        sp = shape.element
        sp.getparent().remove(sp)
    elif shape.shape_type == 13:
        sp = shape.element
        sp.getparent().remove(sp)

px = [0.5, 2.4, 4.3, 6.2, 8.1]
py1 = 2.5
py2 = 4.5

add_box(slide_pipe, px[0], py1, 1.6, 0.8, "1. 帧捕获\nFrame Capture", (224, 255, 255), (0, 139, 139))
add_box(slide_pipe, px[1], py1, 1.6, 0.8, "2. 关键点提取\nMediaPipe 21x3D", (224, 255, 255), (0, 139, 139))
add_box(slide_pipe, px[2], py1, 1.6, 0.8, "3. 静态分类\nKeras 24类", (224, 255, 255), (0, 139, 139))
add_box(slide_pipe, px[3], py1, 1.6, 0.8, "4. 动态追踪\nJ/Z 字母判定", (224, 255, 255), (0, 139, 139))
add_box(slide_pipe, px[4], py1, 1.6, 0.8, "5. 稳定帧确认\n连续15帧机制", (224, 255, 255), (0, 139, 139))

add_box(slide_pipe, px[4], py2, 1.6, 0.8, "6. 运动门控\n防抖与位移过滤", (255, 240, 245), (219, 112, 147))
add_box(slide_pipe, px[3], py2, 1.6, 0.8, "7. 单词提交\nTimeout/空格控制", (255, 240, 245), (219, 112, 147))
add_box(slide_pipe, px[2], py2, 1.6, 0.8, "8. 拼写纠错\nASL混淆字母替换", (255, 250, 205), (184, 134, 11))
add_box(slide_pipe, px[1], py2, 1.6, 0.8, "9. 快捷词展开\n用户级字典映射", (255, 250, 205), (184, 134, 11))
add_box(slide_pipe, px[0], py2, 1.6, 0.8, "10. LLM格式化\n自然英语流利输出", (240, 255, 240), (46, 139, 87))

for i in range(4): add_connector(slide_pipe, px[i]+1.6, py1+0.4, px[i+1], py1+0.4)
add_connector(slide_pipe, px[4]+0.8, py1+0.8, px[4]+0.8, py2)
for i in range(4, 0, -1): add_connector(slide_pipe, px[i], py2+0.4, px[i-1]+1.6, py2+0.4)

prs.save(r'C:\Users\20229\Desktop\ASL_Pro_Studio_Presentation_V3_Native.pptx')
print("PPTX native drawings updated!")
