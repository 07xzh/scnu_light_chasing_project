import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

CLR_BG = RGBColor(0x0D, 0x11, 0x17)
CLR_ACCENT = RGBColor(0x00, 0x7A, 0xFF)
CLR_GREEN = RGBColor(0x10, 0xB9, 0x81)
CLR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CLR_GRAY = RGBColor(0x8B, 0x94, 0xA5)
CLR_PURPLE = RGBColor(0x8B, 0x5C, 0xF6)


def set_bg(slide, color=CLR_BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, text, top=Inches(0.8), size=Pt(40), color=CLR_WHITE):
    txBox = slide.shapes.add_textbox(Inches(1), top, Inches(11), Inches(1))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = size
    p.font.color.rgb = color
    p.font.bold = True
    return tf


def add_subtitle(slide, text, top, size=Pt(20), color=CLR_GRAY):
    txBox = slide.shapes.add_textbox(Inches(1), top, Inches(11), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = size
    p.font.color.rgb = color
    return tf


def add_body(slide, text, top, left=Inches(1), width=Inches(11), size=Pt(18), color=CLR_GRAY):
    txBox = slide.shapes.add_textbox(left, top, width, Inches(5))
    tf = txBox.text_frame
    tf.word_wrap = True
    lines = text.strip().split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = size
        p.font.color.rgb = color
        p.space_after = Pt(6)
    return tf


def add_accent_line(slide, top):
    line = slide.shapes.add_shape(1, Inches(1), top, Inches(0.8), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = CLR_ACCENT
    line.line.fill.background()


# ============================================================
# 第1页：封面
# ============================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(sl)
add_accent_line(sl, Inches(2.5))
add_title(sl, 'ASL Pro Studio', Inches(2.7), size=Pt(56))
add_subtitle(sl, '实时手语识别系统 — 人脸认证 · 个性化快捷词 · LLM 智能格式化', Inches(3.5), size=Pt(24), color=CLR_WHITE)
add_subtitle(sl, 'American Sign Language Recognition with Face ID & Personalized Shortcuts', Inches(4.0), size=Pt(14))
add_subtitle(sl, '2026', Inches(5.5), size=Pt(14))

# ============================================================
# 第2页：项目概述
# ============================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(sl)
add_accent_line(sl, Inches(0.6))
add_title(sl, '项目概述', Inches(0.8))

body = """本项目构建了一个基于计算机视觉的实时美国手语 (ASL) 识别系统，融合人脸识别认证、
个性化快捷词管理与大语言模型智能格式化。系统通过摄像头实时捕捉手部关键点，利用
深度学习模型识别 ASL 字母手势，并自动合成为自然语句。

核心能力：
  ▸ 实时手语识别：24 个静态字母 (A-Y) + J/Z 动态手势追踪
  ▸ Face ID 认证：DeepFace 人脸注册登录 + 深度传感器活体检测
  ▸ 个性化快捷词：登录后可自定义 A-Z 快捷词映射，单字母展开为完整短语
  ▸ LLM 语句格式化：DeepSeek-V4 将拼写单词合成为自然英语句子
  ▸ 多摄像头支持：兼容 Orbbec Femto 深度相机与常规 USB 摄像头"""
add_body(sl, body, Inches(1.8))

# ============================================================
# 第3页：系统架构
# ============================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(sl)
add_accent_line(sl, Inches(0.6))
add_title(sl, '系统架构', Inches(0.8))

body = """┌─────────────────────── 前端 (index.html) ────────────────────────┐
│  实时视频流  │  手势捕获面板  │  句子翻译区  │  侧边栏控制面板     │
│  Canvas 渲染  │  字母/置信度   │  LLM 输出  │  参数调节 + Face ID  │
└─────────────────────┬── WebSocket ──┬────────────────────────────┘
                      │              │
┌───────────────────── 后端 (FastAPI) ───────────────────────────────┐
│                                                                     │
│  ┌────────────┐  ┌─────────────┐  ┌──────────────┐                │
│  │ 摄像头线程  │  │ WS 命令处理  │  │  状态推送    │                │
│  │ (66fps)    │  │ (async)     │  │  (40Hz)     │                │
│  └──────┬─────┘  └──────┬──────┘  └──────┬───────┘                │
│         │               │                │                         │
│  ┌──────┴──────┐  ┌─────┴──────┐  ┌──────┴───────┐               │
│  │ 手势识别管线 │  │submit_word │  │current_state │               │
│  │ MediaPipe   │  │shortcut扩展│  │frame_queue   │               │
│  │ + Keras     │  │login/reg   │  │auth_queue    │               │
│  └─────────────┘  └──────┬─────┘  └──────────────┘               │
│                           │                                        │
│  ┌────────────────────────┼──────────────────────────┐             │
│  │ 拼写纠错  快捷词管理  LLM(DeepSeek)  人脸认证     │             │
│  │ face_auth.py  face_db/  shortcuts.json           │             │
│  └──────────────────────────────────────────────────┘             │
└─────────────────────────────────────────────────────────────────────┘"""
add_body(sl, body, Inches(1.8), size=Pt(12))

# ============================================================
# 第4页：核心识别流水线
# ============================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(sl)
add_accent_line(sl, Inches(0.6))
add_title(sl, '核心识别流水线', Inches(0.8))
body = """帧捕获 -> 手部关键点检测 (MediaPipe, 21点x3D) -> 静态字母分类 (Keras, 24类)
  -> 动态 J/Z 追踪 (方向序列匹配) -> 稳定帧确认 (15帧) -> 运动门控 (21点位移检测)
  -> 字母累积 -> 单词提交 -> 拼写纠错 (pyspellchecker + ASL混淆字母替换)
  -> 快捷词展开 (模式开启+已登录) -> LLM 格式化 (DeepSeek-V4-Flash)

自动超时机制：90帧无手部检测 -> 自动提交单词。8秒无输入 -> 开启新句。

稳定帧确认细节：
  ▸ 同一字母连续15帧 -> 触发提交
  ▸ 运动门控：21个关键点位移超过阈值 -> 重新计数 (防止手势切换误触发)
  ▸ 抖动容错：5帧内字母抖动 -> 自动收敛到新字母

动态J/Z追踪：
  ▸ I/D 稳定触发追踪模式 (45帧观察窗)
  ▸ J: 小指方向序列 S->W/NW (J钩形)
  ▸ Z: 食指方向序列 E->SW->E (Z折线)
  ▸ 追踪失败回退为 I/D"""
add_body(sl, body, Inches(1.8))

# ============================================================
# 第5页：Face ID 认证
# ============================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(sl)
add_accent_line(sl, Inches(0.6))
add_title(sl, 'Face ID 认证系统', Inches(0.8))
body = """人脸检测: MTCNN 级联卷积网络
特征提取: DeepFace + Facenet (128维 embedding 向量)
相似度匹配: 余弦距离 (Cosine Distance, 阈值 0.40)
活体检测: Orbbec 深度传感器 -> 人脸区域深度 STD > 5mm 确认真人
登录投票: 多帧多数表决 + 平均距离平票裁决
光照归一化: CLAHE 自适应直方图均衡化 (LAB 色彩空间 L 通道)
注册融合: 5帧 embedding 平均 -> L2归一化存储

安全机制：
  ▸ 距离检测：30-120cm 有效范围，超出拒绝
  ▸ 无效框过滤：人脸框覆盖全画面 -> 拒识 (误检测)
  ▸ 活体防伪：平面照片/屏幕攻击 -> STD < 5mm -> 拒绝
  ▸ 重试机制：登录最多5次重试，注册最多12次"""
add_body(sl, body, Inches(1.8))

# ============================================================
# 第6页：个性化快捷词系统
# ============================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(sl)
add_accent_line(sl, Inches(0.6))
add_title(sl, '个性化快捷词系统', Inches(0.8))
body = """触发方式: 集成于 submit_word 管线 (空格/自动超时/点击 Chip 均可触发)
编辑界面: 26 字母 A-Z 动态输入网格 (Face ID 登录后才可见)
数据存储: shortcuts.json (用户级持久化, threading.Lock 并发安全)
默认词库: A=Hello, D=Doctor, H=Help (完全可自定义)
LLM 联动: 快捷词展开后仍走 LLM 管道进行自然语句格式化

用户操作流程：
  人脸注册 -> 人脸登录 -> 获取快捷词 -> 开启模式 -> 手势/Chip触发 -> 展开短语 -> LLM成句

前端组件：
  ▸ 侧边栏 Toggle 开关 + 编辑按钮
  ▸ 顶部导航「快捷词管理」标签页
  ▸ 工具栏下方 Chip 快捷面板 (显示已配置的快捷词)

技术细节：
  ▸ submit_word 内优先检查 shortcut_mode_active + current_logged_in_user
  ▸ 命中快捷词 -> 跳过拼写纠错 -> 直接展开 -> LLM 格式化
  ▸ 未命中 -> 走正常拼写纠错流程
  ▸ Chip 点击 -> shortcut_expand WebSocket 命令 -> 后端同步展开"""
add_body(sl, body, Inches(1.8))

# ============================================================
# 第7页：技术栈
# ============================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(sl)
add_accent_line(sl, Inches(0.6))
add_title(sl, '技术栈', Inches(0.8))

body_left = """后端技术栈：
  ▸ Python 3.13 + FastAPI
  ▸ WebSocket 双向通信
  ▸ MediaPipe Hand Landmarker
  ▸ TensorFlow / Keras (model5.keras)
  ▸ DeepFace (Facenet + MTCNN)
  ▸ OpenCV (CLAHE, 图像处理)
  ▸ Orbbec SDK (深度相机)
  ▸ pyspellchecker (拼写纠错)
  ▸ DeepSeek API (OpenAI SDK)
  ▸ threading + asyncio 混合架构"""
add_body(sl, body_left, Inches(1.8), left=Inches(1), width=Inches(5.5))

body_right = """前端技术栈：
  ▸ 纯 HTML/CSS/JS 单页应用
  ▸ Canvas 视频渲染 + 骨架叠加
  ▸ WebSocket 40Hz 状态同步
  ▸ Glass Frost UI 设计风格
  ▸ 响应式 Flexbox 布局

关键参数：
  ▸ STABLE_FRAMES: 15 (可调 5-60)
  ▸ DYNAMIC_WINDOW: 45 (可调 15-90)
  ▸ SENTENCE_TIMEOUT: 8s (可调 2-15)
  ▸ WORD_TIMEOUT: 90 frames
  ▸ MIN_CONFIDENCE: 0.80
  ▸ SIMILARITY_THRESHOLD: 0.40
  ▸ LIVENESS_STD_THRESHOLD: 5mm"""
add_body(sl, body_right, Inches(1.8), left=Inches(7), width=Inches(5.5))

# ============================================================
# 第8页：关键问题与解决方案
# ============================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(sl)
add_accent_line(sl, Inches(0.6))
add_title(sl, '关键问题与解决方案', Inches(0.8))
body = """1. Python 闭包作用域陷阱
   问题：send_loop 嵌套函数赋值模块级全局变量时缺少 global 声明，导致写入的是局部变量
   解决：send_loop 内显式添加 global current_logged_in_user, shortcut_mode_active

2. 登录匹配算法：最低距离制 vs 投票制
   问题：原代码取3帧中距离最低的一帧作为匹配结果，一次偶然低距离即可翻盘
   解决：改为 Counter 多帧投票制，票多者胜；平票时比较平均余弦距离

3. 前端 stale 帧覆盖登录状态
   问题：current_state 帧 (logged_in_user: null) 在 auth_result 后到达，覆盖 _loggedInUser
   解决：result handler 增加 null 防护 — 已有用户名时不允许 null 覆盖

4. 光照环境差异导致人脸识别率下降
   问题：宿舍暗光 vs 实验室亮光 -> Facenet embedding 在光照维度上偏移超过类间差异
   解决：CLAHE 自适应直方图均衡化 (LAB 空间 L 通道, clipLimit=2.0, 8x8 tile)

5. 快捷词与登录状态不当耦合
   问题：_loggedInUser 设置依赖于 data.shortcuts 字段存在 -> 任何原因导致字段缺失则登录失败
   解决：解耦为两个独立分支：_loggedInUser 无条件设置，_userShortcuts 按需设置"""
add_body(sl, body, Inches(1.8))

# ============================================================
# 第9页：演示流程
# ============================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(sl)
add_accent_line(sl, Inches(0.6))
add_title(sl, '演示流程', Inches(0.8))
body = """1. 启动服务器：uvicorn backend.main:app --port 8765
2. 浏览器打开：http://localhost:8765
3. 人脸注册：输入用户名 -> 10秒倒计时 -> 自动多角度采集 -> 5帧融合注册
4. 人脸登录：3秒倒计时 -> DeepFace 匹配 -> 投票制确认 -> 显示 Welcome
5. 开启快捷词模式 (Toggle 开关) -> Chip 快捷面板出现
6. 编辑快捷词：点击「快捷词管理」-> 修改 A-Z 映射 -> 保存
7. 手语输入：面对摄像头比划字母 -> 稳定帧确认 -> 单词拼写 -> LLM 成句
8. 快捷词触发：比划 A -> 自动展开为 "Hello"
9. 点击 Chip：鼠标点击 "D: Doctor" -> 实时追加到句子

预期效果：
  ▸ 实验室环境：登录准确率 95%+，活体防伪有效
  ▸ 宿舍环境：CLAHE 优化后准确率明显提升
  ▸ 快捷词：手势/Chip 双击触发展开 + LLM 智能格式化"""
add_body(sl, body, Inches(1.8))

# ============================================================
# 第10页：未来展望
# ============================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(sl)
add_accent_line(sl, Inches(0.6))
add_title(sl, '未来展望', Inches(0.8))
body = """▸ 多光照/多角度注册融合：注册时采集不同环境样本 -> 平均 embedding 提高鲁棒性
▸ 离线模型推理：TensorFlow Lite 替代 Keras -> 降低 GPU/内存依赖
▸ 中文手语支持：CSL 数据集 + 新分类模型 -> 面向国内应用场景
▸ 移动端适配：WebRTC 视频流 + React Native 前端
▸ 连续手语识别：LSTM / Transformer 序列模型 -> 词组/语句级识别
▸ 语音输出集成：TTS 文字转语音 -> 无障碍沟通闭环
▸ 云端部署：Docker 容器化 + Azure App Service -> 公共访问
▸ 历史记录：识别日志持久化存储 -> 统计分析 + 模型迭代"""
add_body(sl, body, Inches(1.8))

# ============================================================
# 第11页：致谢
# ============================================================
sl = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(sl)
add_accent_line(sl, Inches(2.8))
add_title(sl, '谢谢', Inches(3.0), size=Pt(56))
add_subtitle(sl, 'ASL Pro Studio — 实时手语识别系统', Inches(3.8), size=Pt(24), color=CLR_WHITE)
add_subtitle(sl, 'Face ID · 个性化快捷词 · LLM 智能格式化 · 活体检测', Inches(4.3), size=Pt(14))

# 保存
output_path = os.path.join(os.path.expanduser('~'), 'Desktop', 'ASL_Pro_Studio_Presentation.pptx')
prs.save(output_path)
print(f'PPT saved to: {output_path}')
