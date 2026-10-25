import os
import pptx
from pptx.util import Inches
import docx
from docx.shared import Pt, Inches as DocxInches
import re

arch_img = os.path.join(os.path.dirname(__file__), "..", "docs", "arch_pretty.png")
pipe_img = os.path.join(os.path.dirname(__file__), "..", "docs", "pipe_pretty.png")

# UPDATE PPTX
ppt_path = r'C:\Users\20229\Desktop\ASL_Pro_Studio_Presentation_V2.pptx'
if os.path.exists(ppt_path) and os.path.exists(arch_img):
    try:
        prs = pptx.Presentation(ppt_path)
        # Slide 3: Architecture
        if len(prs.slides) > 2:
            for shape in list(prs.slides[2].shapes):
                if shape.has_text_frame and '系统架构' not in shape.text:
                    sp = shape.element
                    sp.getparent().remove(sp)
            prs.slides[2].shapes.add_picture(arch_img, Inches(0.5), Inches(1.5), width=Inches(8))
        
        # Slide 4: Pipeline
        if len(prs.slides) > 3:
            for shape in list(prs.slides[3].shapes):
                if shape.has_text_frame and '核心识别流水线' not in shape.text:
                    sp = shape.element
                    sp.getparent().remove(sp)
            prs.slides[3].shapes.add_picture(pipe_img, Inches(0.5), Inches(1.5), width=Inches(8))
        
        prs.save(ppt_path)
        print('PPTX updated with images.')
    except Exception as e:
        print('Error updating PPTX:', e)

# GENERATE DOCX from MD
md_path = r'C:\Users\20229\Desktop\ASL_Pro_Studio_项目总结.md'
docx_path = r'C:\Users\20229\Desktop\ASL_Pro_Studio_项目总结.docx'

if os.path.exists(md_path):
    doc = docx.Document()
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_arch_code = False
    for line in lines:
        line = line.strip()
        if not line: continue
        
        if line.startswith('```'):
            if in_arch_code:
                in_arch_code = False
            else:
                in_arch_code = True
            continue
        
        if in_arch_code:
            if 'Canvas' in line and not getattr(doc, '_arch_inserted', False) and os.path.exists(arch_img):
                doc.add_picture(arch_img, width=DocxInches(6))
                doc._arch_inserted = True
            continue
            
        if line.startswith('# '):
            doc.add_heading(line[2:], 0)
        elif line.startswith('## '):
            doc.add_heading(line[3:], 1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], 2)
        elif line.startswith('#### '):
            doc.add_heading(line[5:], 3)
        elif line.startswith('- ') or line.startswith('* '):
            doc.add_paragraph(line[2:], style='List Bullet')
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                elif part.startswith('`') and part.endswith('`'):
                    p.add_run(part).font.name = 'Courier New'
                else:
                    p.add_run(part)

    doc.save(docx_path)
    print('DOCX generated successfully.')
